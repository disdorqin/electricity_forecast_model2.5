#!/usr/bin/env python
"""
国网PMOS爬虫 — 主入口

爬取 → MySQL → 本地文件

用法:
  python scripts/crawler/run_crawler.py                  # 爬取昨天
  python scripts/crawler/run_crawler.py --date 2024-01-15
  python scripts/crawler/run_crawler.py --start 2024-01-01 --end 2024-01-15
  python scripts/crawler/run_crawler.py --init-db        # 建表
  python scripts/crawler/run_crawler.py --no-db          # 仅存文件
  python scripts/crawler/run_crawler.py --no-file        # 仅存DB
"""

from __future__ import annotations

import argparse
import json
import logging
import os
import sys
import time
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Optional

import pandas as pd
from dotenv import load_dotenv

# ---------------------------------------------------------------------------
#  运行时检测 & 路径
# ---------------------------------------------------------------------------

_FROZEN = getattr(sys, "frozen", False)  # PyInstaller .exe 模式

if _FROZEN:
    # .exe 模式：路径相对于可执行文件所在目录
    BASE_DIR = Path(sys.executable).parent.resolve()
    # 在 .exe 中，import 走 PyInstaller 内部 loader，不需要 sys.path 补丁
else:
    # Python 脚本模式：确保项目根在 sys.path 中
    _BASE_DIR = Path(__file__).resolve().parents[2]
    if str(_BASE_DIR) not in sys.path:
        sys.path.insert(0, str(_BASE_DIR))
    BASE_DIR = _BASE_DIR

from scripts.crawler.crawl import PmosCrawler, parse_number, period_no_from_time

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S",
)
logger = logging.getLogger("run_crawler")

# ---------------------------------------------------------------------------
#  Paths
# ---------------------------------------------------------------------------

CRAWLER_DIR = BASE_DIR if _FROZEN else BASE_DIR / "scripts" / "crawler"
OUTPUT_DIR = BASE_DIR / "output"
MIGRATION_SQL = BASE_DIR / "scripts" / "db_migrations" / "001_create_epf_unit_data_96.sql"

# ---------------------------------------------------------------------------
#  市场特征字段映射 (爬虫列名 → DB 列名)
# ---------------------------------------------------------------------------

MARKET_FIELD_MAP: dict[str, str] = {
    "systemload": "actual_direct_load",
    "dfdcload": "actual_local_plant",
    "excload": "actual_tie_line",
    "fdload": "actual_wind",
    "gfload": "actual_solar",
    "sytsjz": "actual_nuclear",
    "selfunit": "actual_self_owned",
    "syjzzj": "actual_test_unit",
}

# ===================================================================
#  配置加载
# ===================================================================


def load_config() -> dict:
    """读取爬虫配置文件 (config.json)"""
    config_path = CRAWLER_DIR / "config.json"
    if not config_path.exists():
        logger.error("配置文件不存在: %s", config_path)
        logger.error("请复制 config.example.json 为 config.json 并填入 Cookie 和机组ID")
        sys.exit(1)

    with open(config_path, "r", encoding="utf-8") as f:
        cfg: dict = json.load(f)

    if not cfg.get("cookie", "").strip():
        logger.error("config.json 中 cookie 为空，请填入有效 Cookie")
        sys.exit(1)
    if not cfg.get("unit_id", "").strip():
        logger.error("config.json 中 unit_id 为空，请填入机组ID")
        sys.exit(1)

    return cfg


def load_db_config() -> dict:
    """读取数据库配置 (从 .env)"""
    load_dotenv(dotenv_path=BASE_DIR / ".env", override=False)

    def _env(key: str) -> str:
        return os.getenv(key, "").strip().strip("\"'")

    cfg = {
        "host": _env("DB_HOST"),
        "port": int(_env("DB_PORT") or "3306"),
        "user": _env("DB_USER"),
        "password": _env("DB_PWD"),
        "database": _env("DB") or _env("DB_NAME"),
    }
    return cfg


# ===================================================================
#  数据库操作
# ===================================================================


def get_db(cfg: dict):
    """创建数据库连接"""
    import pymysql

    return pymysql.connect(
        host=cfg["host"],
        port=cfg["port"],
        user=cfg["user"],
        password=cfg["password"],
        database=cfg["database"],
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
        connect_timeout=10,
    )


# 建表 SQL（无需外部文件，.exe 也可执行）
CREATE_TABLE_SQL = """
CREATE TABLE IF NOT EXISTS `epf_unit_data_96` (
    `id`            BIGINT        NOT NULL AUTO_INCREMENT  COMMENT '主键ID',
    `market_date`   DATE          NOT NULL                 COMMENT '市场日期',
    `period_no`     INT           NOT NULL                 COMMENT '96点序号: 1-96',
    `data_time`     DATETIME      NOT NULL                 COMMENT '完整时刻(区间结束时间)',
    `unit_id`       VARCHAR(64)   NOT NULL                 COMMENT '机组ID',
    `da_cq_price`   DECIMAL(14,4) DEFAULT NULL             COMMENT '日前出清价格(元/MWh)',
    `da_power`      DECIMAL(14,4) DEFAULT NULL             COMMENT '日前出力(MW)',
    `da_energy`     DECIMAL(14,4) DEFAULT NULL             COMMENT '日前电量(MWh)',
    `da_status`     VARCHAR(20)   DEFAULT NULL             COMMENT '日前开机状态',
    `rt_cq_price`   DECIMAL(14,4) DEFAULT NULL             COMMENT '实时出清价格(元/MWh)',
    `rt_power`      DECIMAL(14,4) DEFAULT NULL             COMMENT '实时出力(MW)',
    `rt_energy`     DECIMAL(14,4) DEFAULT NULL             COMMENT '实时电量(MWh)',
    `rt_status`     VARCHAR(20)   DEFAULT NULL             COMMENT '实时开机状态',
    `create_time`   DATETIME      DEFAULT CURRENT_TIMESTAMP COMMENT '创建时间',
    `update_time`   DATETIME      DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP COMMENT '更新时间',
    PRIMARY KEY (`id`),
    UNIQUE KEY `uk_date_period_unit` (`market_date`, `period_no`, `unit_id`),
    KEY `idx_data_time` (`data_time`),
    KEY `idx_market_date` (`market_date`),
    KEY `idx_unit_id` (`unit_id`)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_bin COMMENT='机组级电力市场96点数据(15分钟粒度)';
"""


def init_database_tables(db_cfg: dict) -> bool:
    """初始化数据库表"""
    # 优先读取外部 SQL 文件（方便修改），否则使用内嵌 SQL
    if MIGRATION_SQL.exists():
        sql = MIGRATION_SQL.read_text(encoding="utf-8")
    else:
        sql = CREATE_TABLE_SQL

    statements = [s.strip() for s in sql.split(";") if s.strip()]

    conn = get_db(db_cfg)
    try:
        with conn.cursor() as cur:
            for stmt in statements:
                if stmt.upper().startswith("CREATE") or stmt.upper().startswith("ALTER"):
                    logger.info("执行: %s ...", stmt[:60])
                    cur.execute(stmt)
        conn.commit()
        logger.info("[OK] 数据库表初始化完成")
        return True
    except Exception as e:
        logger.error("建表失败: %s", e)
        conn.rollback()
        return False
    finally:
        conn.close()


def upsert_market_overview(
    conn, market_date: str, rows: list[dict[str, Any]]
) -> int:
    """将市场特征数据 upsert 到 epf_market_data_96"""
    import pymysql

    field_map = MARKET_FIELD_MAP
    db_cols = ["market_date", "period_no", "data_time"] + list(field_map.values())
    placeholders = ", ".join(["%s"] * len(db_cols))
    update_parts = ", ".join([f"{c}=VALUES({c})" for c in field_map.values()])

    sql = (
        f"INSERT INTO epf_market_data_96 ({', '.join(db_cols)}) "
        f"VALUES ({placeholders}) "
        f"ON DUPLICATE KEY UPDATE {update_parts}"
    )

    dt_base = datetime.strptime(market_date, "%Y-%m-%d")
    count = 0
    with conn.cursor() as cur:
        for row in rows:
            period_label = row.get("Periodid", "")
            if not period_label:
                continue
            pno = period_no_from_time(period_label)
            data_time = dt_base + timedelta(minutes=pno * 15)
            vals = [market_date, pno, data_time]
            for crawler_col in field_map:
                vals.append(parse_number(row.get(crawler_col)))
            try:
                cur.execute(sql, vals)
                count += 1
            except pymysql.err.IntegrityError as e:
                logger.warning("跳过 period=%d: %s", pno, e)
    conn.commit()
    logger.info("market_overview: %d periods upserted", count)
    return count


def upsert_unit_data(
    conn,
    market_date: str,
    unit_id: str,
    da_rows: list[dict[str, Any]],
    rt_rows: list[dict[str, Any]],
) -> int:
    """将机组日前/实时数据 upsert 到 epf_unit_data_96"""
    import pymysql

    # 按 period_label 索引 DA/RT 数据
    da_index: dict[str, dict] = {r.get("periodid", ""): r for r in da_rows}
    rt_index: dict[str, dict] = {r.get("periodid", ""): r for r in rt_rows}

    all_labels = sorted(set(da_index.keys()) | set(rt_index.keys()))
    if not all_labels:
        return 0

    db_cols = [
        "market_date", "period_no", "data_time", "unit_id",
        "da_cq_price", "da_power", "da_energy", "da_status",
        "rt_cq_price", "rt_power", "rt_energy", "rt_status",
    ]
    placeholders = ", ".join(["%s"] * len(db_cols))
    update_cols = [c for c in db_cols if c not in ("market_date", "period_no", "data_time", "unit_id")]
    update_parts = ", ".join([f"{c}=VALUES({c})" for c in update_cols])

    sql = (
        f"INSERT INTO epf_unit_data_96 ({', '.join(db_cols)}) "
        f"VALUES ({placeholders}) "
        f"ON DUPLICATE KEY UPDATE {update_parts}"
    )

    dt_base = datetime.strptime(market_date, "%Y-%m-%d")
    count = 0
    with conn.cursor() as cur:
        for label in all_labels:
            pno = period_no_from_time(label)
            data_time = dt_base + timedelta(minutes=pno * 15)
            da = da_index.get(label, {})
            rt = rt_index.get(label, {})

            vals = [
                market_date, pno, data_time, unit_id,
                parse_number(da.get("cqPrice")),
                parse_number(da.get("power")),
                parse_number(da.get("energy")),
                da.get("kt") or None,
                parse_number(rt.get("cqPrice")),
                parse_number(rt.get("power")),
                parse_number(rt.get("energy")),
                rt.get("kt") or None,
            ]
            try:
                cur.execute(sql, vals)
                count += 1
            except pymysql.err.IntegrityError as e:
                logger.warning("跳过 %s period=%d: %s", label, pno, e)
    conn.commit()
    logger.info("unit_data: %d periods upserted", count)
    return count


# ===================================================================
#  本地文件更新
# ===================================================================


def _sanitize_sheet_name(name: str, max_len: int = 31) -> str:
    """Excel sheet name: max 31 chars, no []:*?/\\"""
    safe = "".join(c if c not in "[]:*?/\\" else "_" for c in name)
    return safe[:max_len]


def update_local_files(
    date_str: str,
    unit_id: str,
    da_rows: list[dict],
    rt_rows: list[dict],
    market_rows: list[dict],
) -> None:
    """将爬取数据保存到本地 Excel 备份文件"""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    filename = f"山东电力数据_{date_str}.xlsx"
    filepath = OUTPUT_DIR / filename

    data_map = {
        "日前电价_明细": da_rows,
        "实时电价_明细": rt_rows,
        "市场特征信息": market_rows,
    }

    # 转中文列名
    col_labels = {
        "periodid": "时刻", "cqPrice": "出清价格(元/MWh)", "power": "出力(MW)",
        "energy": "电量(MWh)", "bq": "备注", "kt": "开机状态",
        "Periodid": "时刻", "systemload": "直调负荷(MW)", "dfdcload": "地方电厂发电总加(MW)",
        "excload": "联络线受电负荷(MW)", "fdload": "风电总加(MW)", "gfload": "光伏总加(MW)",
        "sytsjz": "非市场化核电总加(MW)", "selfunit": "自备机组总加(MW)", "syjzzj": "试验机组总加(MW)",
    }

    from openpyxl import Workbook

    wb = Workbook()
    first_sheet = True

    for sheet_name, rows in data_map.items():
        if not rows:
            continue
        safe_name = _sanitize_sheet_name(sheet_name)
        ws = wb.active if first_sheet else wb.create_sheet()
        ws.title = safe_name
        first_sheet = False

        headers = list(rows[0].keys())
        display_headers = [col_labels.get(h, h) for h in headers]
        ws.append(display_headers)
        for row in rows:
            ws.append([str(row.get(h, "")) for h in headers])

    if not first_sheet:
        wb.save(filepath)
        logger.info("本地文件已保存: %s", filepath)
    else:
        logger.warning("无数据，未保存本地文件")


# ===================================================================
#  日期解析
# ===================================================================


def parse_dates(args) -> list[str]:
    """解析 CLI 日期参数"""
    if args.date:
        return [args.date]
    if args.start and args.end:
        start = datetime.strptime(args.start, "%Y-%m-%d")
        end = datetime.strptime(args.end, "%Y-%m-%d")
        dates = []
        cur = start
        while cur <= end:
            dates.append(cur.strftime("%Y-%m-%d"))
            cur += timedelta(days=1)
        return dates
    # 默认昨天
    yesterday = date.today() - timedelta(days=1)
    return [yesterday.strftime("%Y-%m-%d")]


# ===================================================================
#  主流程
# ===================================================================


def main() -> None:
    parser = argparse.ArgumentParser(description="国网PMOS爬虫")
    parser.add_argument("--date", help="爬取指定日期 (YYYY-MM-DD)")
    parser.add_argument("--start", help="开始日期 (YYYY-MM-DD)")
    parser.add_argument("--end", help="结束日期 (YYYY-MM-DD)")
    parser.add_argument("--init-db", action="store_true", help="初始化数据库表")
    parser.add_argument("--no-db", action="store_true", help="不写入数据库")
    parser.add_argument("--no-file", action="store_true", help="不保存本地文件")
    args = parser.parse_args()

    print("=" * 55)
    print("  山东电力市场数据爬取 — 项目集成版")
    print("=" * 55)

    # 1. 加载配置
    config = load_config()
    cookie = config["cookie"]
    unit_id = config["unit_id"]
    base_url = config.get("base_url", "https://pmos.sd.sgcc.com.cn:18080/trade")

    # 2. 数据库
    db_cfg = load_db_config()
    db_ok = all([db_cfg.get("host"), db_cfg.get("database"), db_cfg.get("user"), db_cfg.get("password")])

    if args.init_db:
        if not db_ok:
            logger.error("数据库配置不完整，无法初始化")
            sys.exit(1)
        logger.info("初始化数据库表...")
        init_database_tables(db_cfg)
        return

    # 3. 日期列表
    dates = parse_dates(args)
    logger.info("待爬取日期: %s", dates)

    # 4. 逐日爬取
    for idx, date_str in enumerate(dates):
        print(f"\n{'─' * 50}")
        print(f"  [{idx + 1}/{len(dates)}] {date_str}")
        print(f"{'─' * 50}")

        spider = PmosCrawler(base_url=base_url, cookie=cookie, unit_id=unit_id)

        # 4a. 认证
        if not spider.fetch_csrf_token():
            logger.warning("CSRF token 获取失败，尝试继续...")

        # 4b. 切换日期
        if not spider.change_date(date_str):
            logger.error("日期切换失败，跳过 %s", date_str)
            continue

        time.sleep(1)

        # 4c. 爬取市场特征
        market_rows: list[dict] = []
        logger.info("爬取市场特征信息...")
        try:
            market_rows = spider.crawl_market_overview()
            logger.info("  → %d 行", len(market_rows))
        except Exception as e:
            logger.error("市场特征爬取失败: %s", e)

        # 4d. 爬取日前电价
        da_rows: list[dict] = []
        logger.info("爬取日前电价明细...")
        try:
            da_rows = spider.crawl_day_ahead()
            logger.info("  → %d 行", len(da_rows))
        except Exception as e:
            logger.error("日前电价爬取失败: %s", e)

        # 4e. 爬取实时电价
        rt_rows: list[dict] = []
        logger.info("爬取实时电价明细...")
        try:
            rt_rows = spider.crawl_realtime()
            logger.info("  → %d 行", len(rt_rows))
        except Exception as e:
            logger.error("实时电价爬取失败: %s", e)

        # 4f. 写入数据库
        if not args.no_db and db_ok:
            try:
                conn = get_db(db_cfg)
                try:
                    if market_rows:
                        upsert_market_overview(conn, date_str, market_rows)
                    if da_rows or rt_rows:
                        upsert_unit_data(conn, date_str, unit_id, da_rows, rt_rows)
                finally:
                    conn.close()
                logger.info("[DB] 写入完成")
            except Exception as e:
                logger.error("数据库写入失败: %s", e)
        elif not args.no_db and not db_ok:
            logger.warning("数据库配置不完整，跳过 DB 写入")

        # 4g. 保存本地文件
        if not args.no_file:
            try:
                update_local_files(date_str, unit_id, da_rows, rt_rows, market_rows)
            except Exception as e:
                logger.error("本地文件保存失败: %s", e)

    print(f"\n{'=' * 55}")
    logger.info("全部完成！")
    print(f"{'=' * 55}")


if __name__ == "__main__":
    main()
